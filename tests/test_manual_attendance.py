"""Staff correcting attendance by hand."""
import io

import openpyxl

from attendqr.extensions import db
from attendqr.models import ActivityLog, Admin, Attendance, Student
from attendqr.services.reports import build_student_attendance_summary
from conftest import add_session, post


def mark(client, class_session, student_id, note='Phone battery died'):
    return post(client, f'/sessions/{class_session.id}/attendance', {'student_id': student_id, 'note': note})


def test_teacher_marks_student_present_with_reason(data, client_for):
    class_session = add_session(data)
    response = mark(client_for(data.teacher), class_session, data.student_id)
    assert response.status_code == 302
    assert response.headers['Location'].endswith(f'/generate_qr?session_id={class_session.id}#qr-attendance')

    entry = Attendance.query.filter_by(session_id=class_session.id, student_id=data.student_id).one()
    teacher_account = Admin.query.filter_by(username=data.teacher).one()
    assert entry.source == 'manual' and entry.is_manual
    assert entry.marked_by_admin_id == teacher_account.id
    assert entry.note == 'Phone battery died'

    log = ActivityLog.query.filter_by(action='attendance_marked_manually').one()
    assert 'Rahul Das (MCA001)' in log.description and 'Phone battery died' in log.description
    assert log.actor_username == data.teacher


def test_reason_is_required(data, client_for):
    class_session = add_session(data)
    response = mark(client_for(data.teacher), class_session, data.student_id, note='   ')
    assert response.status_code == 302
    assert Attendance.query.count() == 0
    page = client_for(data.teacher).get(f'/generate_qr?session_id={class_session.id}')
    assert page.status_code == 200


def test_cannot_mark_student_from_another_class(data, client_for):
    class_session = add_session(data)
    late = Student.query.filter_by(email=data.sem2_student).one()
    mark(client_for(data.teacher), class_session, late.id)
    assert Attendance.query.count() == 0


def test_duplicate_manual_mark_is_ignored(data, client_for):
    class_session = add_session(data, present=[data.student_id])
    response = mark(client_for(data.teacher), class_session, data.student_id)
    assert response.status_code == 302
    entries = Attendance.query.filter_by(session_id=class_session.id).all()
    assert len(entries) == 1 and entries[0].source == 'scan'


def test_only_staff_who_own_the_session_can_correct(data, client_for):
    class_session = add_session(data)
    assert mark(client_for(data.other_teacher), class_session, data.student_id).status_code == 403
    assert mark(client_for(data.student), class_session, data.student_id).status_code == 403
    assert mark(client_for(data.hod), class_session, data.student_id).status_code == 302
    assert Attendance.query.count() == 1


def test_remove_attendance_requires_reason_and_is_logged(data, client_for):
    class_session = add_session(data, present=[data.student_id])
    entry = Attendance.query.one()
    teacher = client_for(data.teacher)
    url = f'/sessions/{class_session.id}/attendance/{entry.id}/remove'

    post(teacher, url, {'note': ''})
    assert Attendance.query.count() == 1

    post(teacher, url, {'note': 'Marked from outside the classroom'})
    assert Attendance.query.count() == 0
    log = ActivityLog.query.filter_by(action='attendance_removed').one()
    assert 'scan entry' in log.description and 'outside the classroom' in log.description


def test_remove_rejects_entry_from_another_session(data, client_for):
    first = add_session(data, present=[data.student_id])
    second = add_session(data)
    entry = Attendance.query.filter_by(session_id=first.id).one()
    response = post(client_for(data.teacher), f'/sessions/{second.id}/attendance/{entry.id}/remove', {'note': 'x'})
    assert response.status_code == 404
    assert db.session.get(Attendance, entry.id) is not None


def test_session_page_offers_only_unmarked_students(data, client_for):
    class_session = add_session(data, present=[data.student_id])
    html = client_for(data.teacher).get(f'/generate_qr?session_id={class_session.id}').get_data(as_text=True)
    picker = html.split('id="manual_student"')[1].split('</select>')[0]
    assert 'Priya Sen' in picker and 'Rahul Das' not in picker and 'Late Joiner' not in picker


def test_manual_entries_are_labelled_in_reports(data, client_for):
    class_session = add_session(data, present=[data.student2_id])
    teacher = client_for(data.teacher)
    mark(teacher, class_session, data.student_id)

    detailed = teacher.get(f'/reports/detailed?department_id={data.mca_id}&semester_id={data.sem1_id}')
    assert '>P*</span>' in detailed.get_data(as_text=True)

    api = teacher.get(f'/api/session_attendance/{class_session.id}').get_json()
    sources = {a['student_name']: (a['source'], a['marked_by']) for a in api['attendances']}
    assert sources == {'Rahul Das': ('manual', data.teacher), 'Priya Sen': ('scan', '')}

    sheet = openpyxl.load_workbook(io.BytesIO(teacher.get(f'/export/{data.mca_id}/{data.sem1_id}/excel').data)).active
    marks = {row[1]: row[3] for row in sheet.iter_rows(min_row=2, max_row=3, values_only=True)}
    assert marks == {'Rahul Das': 'P*', 'Priya Sen': 'P'}
    assert any('P* = present (entered manually' in str(row[0]) for row in sheet.iter_rows(values_only=True))

    pdf = teacher.get(f'/export/{data.mca_id}/{data.sem1_id}/pdf')
    assert pdf.status_code == 200 and pdf.data.startswith(b'%PDF')


def test_manual_mark_makes_a_void_session_count(data, client_for):
    class_session = add_session(data, active=False)  # deactivated before anyone scanned: void
    student = db.session.get(Student, data.student_id)
    dbms = lambda: next(r for r in build_student_attendance_summary(student)['rows'] if r['subject'].id == data.subject_id)  # noqa: E731
    assert dbms()['total'] == 0

    mark(client_for(data.teacher), class_session, data.student_id)
    assert (dbms()['attended'], dbms()['total']) == (1, 1)


def test_student_sees_manual_entries_marked_as_such(data, client_for):
    class_session = add_session(data)
    mark(client_for(data.teacher), class_session, data.student_id)
    assert 'Present (marked by staff)' in client_for(data.student).get('/').get_data(as_text=True)


def test_deleting_staff_account_keeps_their_manual_entries(data, client_for):
    class_session = add_session(data)
    mark(client_for(data.hod), class_session, data.student_id)
    hod = Admin.query.filter_by(username=data.hod).one()
    assert post(client_for(data.principal), f'/delete_admin/{hod.id}').status_code == 302

    db.session.expire_all()
    entry = Attendance.query.one()
    assert entry.source == 'manual' and entry.marked_by_admin_id is None
    page = client_for(data.teacher).get(f'/generate_qr?session_id={class_session.id}').get_data(as_text=True)
    assert 'Deleted account' in page
